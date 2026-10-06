import openpyxl
import pandas as pd
import numpy as np
import os
import sys
import csv

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

datasets = [
    {
        'file_id': 'c_strength',
        'filename': 'C Strenth Reading 1000000.xlsx',
        'filepath': 'data/C Strenth Reading 1000000.xlsx',
        'primary_target': 'Compressive Strength (MPa)'
    },
    {
        'file_id': 'f_strength',
        'filename': 'F Strength Reading 1000000.xlsx',
        'filepath': 'data/F Strength Reading 1000000.xlsx',
        'primary_target': 'Flexural Strength (MPa)'
    },
    {
        'file_id': 't_strength',
        'filename': 'T Strenth Reading 1000000.xlsx',
        'filepath': 'data/T Strenth Reading 1000000.xlsx',
        'primary_target': 'Split Tensile Strength (MPa)'
    }
]

schema_rows = []
audit_data = {}

for d in datasets:
    fname = d['filename']
    fpath = d['filepath']
    wb_formulas = openpyxl.load_workbook(fpath, data_only=False)
    wb_values = openpyxl.load_workbook(fpath, data_only=True)
    
    file_record = {
        'filename': fname,
        'filepath': fpath,
        'file_size_bytes': os.path.getsize(fpath),
        'primary_target': d['primary_target'],
        'sheets': {}
    }
    
    for sname in wb_values.sheetnames:
        sheet_f = wb_formulas[sname]
        sheet_v = wb_values[sname]
        
        # Read raw grid
        raw_v_rows = list(sheet_v.iter_rows(values_only=True))
        raw_f_rows = list(sheet_f.iter_rows(values_only=True))
        
        total_rows = len(raw_v_rows)
        total_cols = sheet_v.max_column if sheet_v.max_column else 0
        
        # Formula detection
        formulas_found = []
        for r_i, r in enumerate(raw_f_rows):
            for c_i, c in enumerate(r):
                if isinstance(c, str) and c.startswith('='):
                    formulas_found.append({
                        'cell': f"{openpyxl.utils.get_column_letter(c_i+1)}{r_i+1}",
                        'formula': c
                    })
        
        # Determine header and clean DataFrame
        # Raw reading with pandas
        df_raw = pd.read_excel(fpath, sheet_name=sname, header=None)
        
        # Identify non-empty rows
        non_empty_indices = [i for i, r in enumerate(raw_v_rows) if any(cell is not None for cell in r)]
        non_empty_count = len(non_empty_indices)
        
        sheet_record = {
            'sheet_name': sname,
            'total_rows': total_rows,
            'total_cols': total_cols,
            'non_empty_rows': non_empty_count,
            'formulas_count': len(formulas_found),
            'sample_formulas': [f['cell'] + ': ' + f['formula'] for f in formulas_found[:5]],
            'raw_duplicates': int(df_raw.duplicated().sum()) if not df_raw.empty else 0,
            'columns': []
        }
        
        # Now parse structured table for this sheet
        # Case by case to get accurate column names and types
        if fname.startswith('C') and sname == 'Comp 6':
            # Row 0: ['Concrete Type', 'Curing Age (Days)', 'Cement', 'Bacteria', 'Bacterial Concentration (cells/mL)', 'Compressive Strength (MPa)', '']
            # Row 1: ['', '', '', '', '', 'N', 'BC']
            df_clean = pd.read_excel(fpath, sheet_name=sname, skiprows=2, header=None)
            df_clean.columns = [
                'Concrete Type', 'Curing Age (Days)', 'Cement', 'Bacteria',
                'Bacterial Concentration (cells/mL)',
                'Compressive Strength - Normal Concrete (MPa)',
                'Compressive Strength - Bacterial Concrete (MPa)'
            ]
            sheet_record['interpretation'] = 'Primary Raw Dataset (Paired N and BC samples across 6 curing ages)'
            sheet_record['data_nature'] = 'Raw experimental sample measurements'
        elif fname.startswith('F') and sname == 'Sheet1':
            df_clean = pd.read_excel(fpath, sheet_name=sname)
            sheet_record['interpretation'] = 'Primary Raw Dataset (Long-format N and BC, note 200 missing values in BC 56/90d)'
            sheet_record['data_nature'] = 'Raw experimental sample measurements (unpivoted long format)'
        elif fname.startswith('T') and sname == 'Sheet1':
            df_clean = pd.read_excel(fpath, sheet_name=sname)
            sheet_record['interpretation'] = 'Primary Raw Dataset (Long-format N and BC, complete 1200 rows)'
            sheet_record['data_nature'] = 'Raw experimental sample measurements (unpivoted long format)'
        elif sname in ['Sheet2', 'Ten 6', 'Sheet1'] and total_cols in [13, 14]:
            # Wide sample matrix: find row with 'Sample No'
            header_row_idx = None
            for idx, r in enumerate(raw_v_rows[:10]):
                if any(isinstance(c, str) and 'Sample No' in c for c in r if c is not None):
                    header_row_idx = idx
                    break
            if header_row_idx is not None:
                df_clean = pd.read_excel(fpath, sheet_name=sname, skiprows=header_row_idx)
                # drop completely empty columns
                df_clean = df_clean.dropna(how='all', axis=1)
            else:
                df_clean = df_raw.copy()
            sheet_record['interpretation'] = 'Matrix Format Dataset (100 physical samples x 12 test conditions + summary statistics)'
            sheet_record['data_nature'] = 'Raw experimental measurements (samples 1..100) followed by calculated/summary statistics'
        elif sname in ['c', 'f', 't']:
            # Header is around row 4 or 5
            header_row_idx = None
            for idx, r in enumerate(raw_v_rows[:10]):
                if any(isinstance(c, str) and c.strip() == 'Days' for c in r if c is not None):
                    header_row_idx = idx
                    break
            if header_row_idx is not None:
                df_clean = pd.read_excel(fpath, sheet_name=sname, skiprows=header_row_idx)
                df_clean = df_clean.dropna(how='all', axis=1).dropna(how='all', axis=0)
            else:
                df_clean = df_raw.copy()
            sheet_record['interpretation'] = f'Multi-Concentration Summary Table for {sname.upper()} Strength across BC3..BC8'
            sheet_record['data_nature'] = 'Calculated / summary values (Averages across curing days and concentrations)'
        elif sname == 'RCPT':
            # Header has 'Bacterial Dose'
            header_row_idx = None
            for idx, r in enumerate(raw_v_rows[:10]):
                if any(isinstance(c, str) and 'Bacterial Dose' in c for c in r if c is not None):
                    header_row_idx = idx
                    break
            if header_row_idx is not None:
                df_clean = pd.read_excel(fpath, sheet_name=sname, skiprows=header_row_idx)
                df_clean = df_clean.dropna(how='all', axis=1).dropna(how='all', axis=0)
            else:
                df_clean = df_raw.copy()
            sheet_record['interpretation'] = 'Durability: Rapid Chloride Permeability Test (RCPT) in Coulombs'
            sheet_record['data_nature'] = 'Calculated / benchmark durability test values'
        elif sname == 'Water Absorption':
            header_row_idx = None
            for idx, r in enumerate(raw_v_rows[:10]):
                if any(isinstance(c, str) and 'Water Absorption' in c for c in r if c is not None):
                    header_row_idx = idx
                    break
            if header_row_idx is not None:
                df_clean = pd.read_excel(fpath, sheet_name=sname, skiprows=header_row_idx)
                df_clean = df_clean.dropna(how='all', axis=1).dropna(how='all', axis=0)
            else:
                df_clean = df_raw.copy()
            sheet_record['interpretation'] = 'Durability: Water Absorption (%) across bacterial concentrations'
            sheet_record['data_nature'] = 'Calculated / experimental water absorption percentages'
        else:
            df_clean = df_raw.copy()
            sheet_record['interpretation'] = f'Supplementary / derived comparison sheet: {sname}'
            sheet_record['data_nature'] = 'Derived / comparison summary values'
            
        sheet_record['clean_shape'] = df_clean.shape
        sheet_record['clean_duplicates'] = int(df_clean.duplicated().sum()) if not df_clean.empty else 0
        
        # Analyze each column in df_clean
        for col_idx, col in enumerate(df_clean.columns):
            col_series = df_clean[col]
            c_name = str(col).strip()
            total_c = len(col_series)
            null_c = int(col_series.isnull().sum())
            non_null_c = total_c - null_c
            missing_pct = round((null_c / total_c * 100), 2) if total_c > 0 else 0.0
            
            # Numeric conversion check
            numeric_series = pd.to_numeric(col_series, errors='coerce')
            is_numeric = numeric_series.notnull().sum() > (non_null_c * 0.7) and non_null_c > 0
            
            # Unique values
            # dropna for unique
            val_counts = col_series.dropna().unique()
            uniq_c = len(val_counts)
            is_const = (uniq_c <= 1) and (non_null_c > 0)
            
            # Target identification
            is_target = False
            target_type = 'N/A'
            c_lower = c_name.lower()
            if 'compressive' in c_lower or (c_name in ['c', '7 N', '7 BC', '14 N', '14 BC', '21 N', '21 BC', '28 N', '28 BC', '56 N', '56 BC', '90 N', '90 BC'] and fname.startswith('C')):
                is_target = True
                target_type = 'Compressive Strength'
            elif 'flexural' in c_lower or (c_name in ['f', '7 N', '7 BC', '14 N', '14 BC', '21 N', '21 BC', '28 N', '28 BC', '56 N', '56 BC', '90 N', '90 BC'] and fname.startswith('F')):
                is_target = True
                target_type = 'Flexural Strength'
            elif 'tensile' in c_lower or (c_name in ['t', '7 N', '7 BC', '14 N', '14 BC', '21 N', '21 BC', '28 N', '28 BC', '56 N', '56 BC', '90 N', '90 BC'] and fname.startswith('T')):
                is_target = True
                target_type = 'Tensile Strength'
            elif 'rcpt' in c_lower:
                is_target = True
                target_type = 'Rapid Chloride Permeability'
            elif 'water absorption' in c_lower:
                is_target = True
                target_type = 'Water Absorption'
            
            # Categorize measurement nature
            if is_const:
                nature = 'Constant Feature'
            elif is_target:
                nature = 'Target Variable'
            elif 'age' in c_lower or 'day' in c_lower:
                nature = 'Input Feature (Curing Age)'
            elif 'sample' in c_lower or 'sr.' in c_lower:
                nature = 'Identifier / Index'
            elif 'concrete' in c_lower or 'mix' in c_lower:
                nature = 'Input Feature (Concrete Type)'
            elif 'dose' in c_lower or 'concentration' in c_lower:
                nature = 'Input Feature (Bacterial Concentration)'
            elif 'bacteria' in c_lower or 'cement' in c_lower:
                nature = 'Input Feature (Categorical)'
            else:
                nature = 'Feature / Reading'

            col_meta = {
                'column_index': col_idx + 1,
                'column_name': c_name,
                'data_type': str(col_series.dtype),
                'total_count': total_c,
                'non_null_count': non_null_c,
                'missing_count': null_c,
                'missing_percentage': missing_pct,
                'unique_count': uniq_c,
                'is_constant': is_const,
                'unique_values_sample': [str(v) for v in val_counts[:6]] if not is_numeric else [],
                'min': float(numeric_series.min()) if is_numeric and numeric_series.notnull().sum() > 0 else None,
                'max': float(numeric_series.max()) if is_numeric and numeric_series.notnull().sum() > 0 else None,
                'mean': float(numeric_series.mean()) if is_numeric and numeric_series.notnull().sum() > 0 else None,
                'median': float(numeric_series.median()) if is_numeric and numeric_series.notnull().sum() > 0 else None,
                'std': float(numeric_series.std()) if is_numeric and numeric_series.notnull().sum() > 1 else None,
                'is_target': is_target,
                'target_type': target_type,
                'measurement_nature': nature
            }
            sheet_record['columns'].append(col_meta)
            
            # Add to schema_rows
            schema_rows.append({
                'dataset_file': fname,
                'sheet_name': sname,
                'column_index': col_idx + 1,
                'column_name': c_name,
                'data_type': str(col_series.dtype),
                'total_count': total_c,
                'non_null_count': non_null_c,
                'missing_count': null_c,
                'missing_percentage': missing_pct,
                'unique_count': uniq_c,
                'is_constant': is_const,
                'min_value': round(col_meta['min'], 4) if col_meta['min'] is not None else '',
                'max_value': round(col_meta['max'], 4) if col_meta['max'] is not None else '',
                'mean_value': round(col_meta['mean'], 4) if col_meta['mean'] is not None else '',
                'median_value': round(col_meta['median'], 4) if col_meta['median'] is not None else '',
                'std_value': round(col_meta['std'], 4) if col_meta['std'] is not None else '',
                'is_target': is_target,
                'target_type': target_type,
                'measurement_nature': nature
            })
            
        file_record['sheets'][sname] = sheet_record
        
    audit_data[d['file_id']] = file_record

# Save reports/dataset_schema.csv
schema_df = pd.DataFrame(schema_rows)
schema_df.to_csv('reports/dataset_schema.csv', index=False, encoding='utf-8')
print("Successfully generated reports/dataset_schema.csv with", len(schema_df), "columns documented.")
