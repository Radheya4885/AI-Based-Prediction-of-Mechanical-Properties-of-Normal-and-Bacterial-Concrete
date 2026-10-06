import openpyxl
import pandas as pd
import numpy as np
import sys
import json

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

files = {
    'C Strenth Reading 1000000.xlsx': 'data/C Strenth Reading 1000000.xlsx',
    'F Strength Reading 1000000.xlsx': 'data/F Strength Reading 1000000.xlsx',
    'T Strenth Reading 1000000.xlsx': 'data/T Strenth Reading 1000000.xlsx'
}

for fname, fpath in files.items():
    print(f"\n=======================================================")
    print(f"FILE: {fname}")
    print(f"=======================================================")
    wb = openpyxl.load_workbook(fpath, data_only=True)
    for sname in wb.sheetnames:
        sheet = wb[sname]
        rows = list(sheet.iter_rows(values_only=True))
        # Find non-empty rows
        non_empty = [(idx+1, r) for idx, r in enumerate(rows) if any(c is not None for c in r)]
        print(f"\n--- Sheet: '{sname}' | Total Excel Rows: {len(rows)} | Non-empty rows: {len(non_empty)} | Max Col: {sheet.max_column} ---")
        for r_idx, r in non_empty[:15]:
            vals = [str(c) if c is not None else "" for c in r]
            # print up to 15 columns
            print(f"  Row {r_idx:3d}: {vals[:15]}")
        if len(non_empty) > 15:
            print(f"  ... ({len(non_empty)-15} more rows) ...")
            # print last 3 rows
            for r_idx, r in non_empty[-3:]:
                vals = [str(c) if c is not None else "" for c in r]
                print(f"  Row {r_idx:3d}: {vals[:15]}")
